"""PowerPoint decks, built and read back with python-pptx.

Unlike docx and xlsx there is no hand-built fixture here. A deck's structure is
mostly layouts and masters that python-pptx supplies from a real template, and
reproducing those by hand would test the reproduction rather than the handler.
What matters instead is the recursion: a chart carries a complete workbook inside
the presentation, and the leak checks look inside that too.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest

pptx_module = pytest.importorskip("pptx", reason="python-pptx is a dev dependency")
pytest.importorskip("lxml")

from freedocs.formats.ooxml.pptx import PptxHandler  # noqa: E402
from freedocs.formats.router import handler_for  # noqa: E402
from freedocs.pipeline import Scrubber, build_registry  # noqa: E402

from .conftest import CANARIES, l1_scrubber  # noqa: E402

#: Well outside a 9144000 EMU wide slide, so the shape is invisible on screen.
OFF_CANVAS_X = -4_000_000


def build_deck(
    path: Path,
    *,
    chart: bool = True,
    off_canvas: bool = True,
    hidden_slide: bool = False,
    alt_text: str | None = None,
) -> Path:
    from pptx.util import Emu, Inches

    presentation = pptx_module.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Clienti riservati"
    slide.placeholders[1].text = f"Mario Rossi - {CANARIES['IT_CODICE_FISCALE']}"

    # Speaker notes: invisible while presenting, and where the blunt version of
    # the sentence goes.
    slide.notes_slide.notes_text_frame.text = f"Nota relatore: IBAN {CANARIES['IBAN']}"

    if off_canvas:
        box = slide.shapes.add_textbox(Emu(OFF_CANVAS_X), Inches(1), Inches(3), Inches(1))
        box.text_frame.text = f"AVS {CANARIES['CH_AHV']}"

    if chart:
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE

        data = CategoryChartData()
        data.categories = ["Mario Rossi", "Anna Bianchi"]
        data.add_series("Importi", (1500, 2300))
        slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(3), Inches(4), Inches(3), data
        )

    if alt_text is not None:
        shape = slide.shapes.add_textbox(Inches(1), Inches(6), Inches(2), Inches(1))
        shape._element._nvXxPr.cNvPr.set("descr", alt_text)

    if hidden_slide:
        extra = presentation.slides.add_slide(presentation.slide_layouts[1])
        extra.shapes.title.text = f"UID {CANARIES['CH_UID']}"
        extra._element.set("show", "0")

    presentation.core_properties.author = "Marco Tribuzio"
    presentation.core_properties.title = "Riservato"
    presentation.save(path)
    return path


def names_in(data: bytes) -> list[str]:
    return sorted(zipfile.ZipFile(io.BytesIO(data)).namelist())


def survives_anywhere(data: bytes, value: str) -> str | None:
    """Search every part, and every part of every embedded package.

    Returns the part it was found in, so a failure says where rather than only
    that. A chart's embedded workbook is a zip inside a zip, so a single-level
    search would report a clean file that is not.
    """
    needle = value.encode()
    archive = zipfile.ZipFile(io.BytesIO(data))
    for name in archive.namelist():
        body = archive.read(name)
        if needle in body:
            return name
        if name.lower().endswith((".xlsx", ".xlsm", ".docx")):
            nested = zipfile.ZipFile(io.BytesIO(body))
            for inner in nested.namelist():
                if needle in nested.read(inner):
                    return f"{name}!{inner}"
    return None


def texts(data: bytes) -> list[str]:
    presentation = pptx_module.Presentation(io.BytesIO(data))
    out = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text:
                out.append(shape.text_frame.text)
        out.append(slide.notes_slide.notes_text_frame.text)
    return out


@pytest.fixture
def source(tmp_path: Path) -> Path:
    return build_deck(tmp_path / "deck.pptx")


@pytest.fixture
def scrubbed(source: Path) -> bytes:
    return l1_scrubber().scrub_file(source).output


# -- routing and fixture sanity ----------------------------------------------


def test_router_picks_the_pptx_handler(source: Path):
    assert isinstance(handler_for(source), PptxHandler)


def test_a_deck_without_a_presentation_part_is_refused(tmp_path: Path):
    path = tmp_path / "vuoto.pptx"
    path.write_bytes(b"PK\x03\x04" + b"\x00" * 64)
    with pytest.raises(Exception):  # noqa: B017 - any failure beats silent success
        handler_for(path)


def test_the_fixture_really_embeds_a_workbook(source: Path):
    """If python-pptx stops embedding one, the recursion tests stop meaning anything."""
    assert any(name.startswith("ppt/embeddings/") for name in names_in(source.read_bytes()))
    assert any("thumbnail" in name for name in names_in(source.read_bytes()))


# -- extraction ---------------------------------------------------------------


def test_speaker_notes_are_extracted(source: Path):
    stream = handler_for(source).extract()
    assert "Nota relatore" in stream.text
    assert CANARIES["IBAN"] in stream.text


def test_an_off_canvas_shape_is_extracted(source: Path):
    """Invisible on screen, fully present in the XML."""
    stream = handler_for(source).extract()
    assert CANARIES["CH_AHV"] in stream.text


def test_chart_cache_labels_are_extracted(source: Path):
    stream = handler_for(source).extract()
    assert "Anna Bianchi" in stream.text


def test_the_embedded_workbook_is_extracted(source: Path):
    """The chart's source workbook is a complete package inside this one."""
    handler = handler_for(source)
    assert handler.nested_handlers(), "no embedded workbook was picked up"
    stream = handler.extract()
    assert any(anchor.part.startswith("xl/") for anchor in stream.anchors)


def test_alt_text_is_extracted(tmp_path: Path):
    """Someone typed it, so it can say anything."""
    path = build_deck(
        tmp_path / "alt.pptx", chart=False, off_canvas=False, alt_text=f"IBAN {CANARIES['IBAN']}"
    )
    stream = handler_for(path).extract()
    assert stream.text.count(CANARIES["IBAN"]) >= 2, "alt text and notes both hold it"


def test_shape_text_is_not_fused_with_the_next_shape_name(source: Path):
    """Without a boundary the detector reads two shapes as one string."""
    stream = handler_for(source).extract()
    assert "Clienti riservatiContent" not in stream.text


# -- output integrity ---------------------------------------------------------


def test_the_scrubbed_deck_still_opens(scrubbed: bytes):
    presentation = pptx_module.Presentation(io.BytesIO(scrubbed))
    assert len(presentation.slides) == 1


def test_the_chart_survives(scrubbed: bytes):
    names = names_in(scrubbed)
    assert any(name.startswith("ppt/charts/") for name in names)
    assert any(name.startswith("ppt/embeddings/") for name in names)


def test_layouts_and_theme_are_untouched(source: Path, scrubbed: bytes):
    original = zipfile.ZipFile(io.BytesIO(source.read_bytes()))
    result = zipfile.ZipFile(io.BytesIO(scrubbed))
    for name in ("ppt/theme/theme1.xml", "ppt/tableStyles.xml"):
        assert result.read(name) == original.read(name), name


def test_visible_non_confidential_text_is_left_alone(scrubbed: bytes):
    assert "Clienti riservati" in texts(scrubbed)


# -- redaction ----------------------------------------------------------------


def test_the_slide_body_is_redacted(scrubbed: bytes):
    assert any("[TAX_ID_1]" in text for text in texts(scrubbed))


def test_the_notes_are_redacted(scrubbed: bytes):
    assert any(text == "Nota relatore: IBAN [IBAN_1]" for text in texts(scrubbed))


def test_the_off_canvas_shape_is_redacted(scrubbed: bytes):
    assert any(text == "AVS [NATIONAL_ID_1]" for text in texts(scrubbed))


# -- leaks --------------------------------------------------------------------


@pytest.mark.leak
@pytest.mark.parametrize(
    "value",
    [CANARIES["IT_CODICE_FISCALE"], CANARIES["IBAN"], CANARIES["CH_AHV"]],
)
def test_no_canary_survives_anywhere(scrubbed: bytes, value: str):
    assert survives_anywhere(scrubbed, value) is None


@pytest.mark.leak
def test_a_canary_on_a_hidden_slide_does_not_survive(tmp_path: Path):
    path = build_deck(tmp_path / "nascosta.pptx", chart=False, hidden_slide=True)
    output = l1_scrubber().scrub_file(path).output
    assert survives_anywhere(output, CANARIES["CH_UID"]) is None


@pytest.mark.leak
def test_the_thumbnail_is_removed(source: Path, scrubbed: bytes):
    assert not any("thumbnail" in name for name in names_in(scrubbed))


@pytest.mark.leak
def test_metadata_does_not_survive(scrubbed: bytes):
    presentation = pptx_module.Presentation(io.BytesIO(scrubbed))
    assert not presentation.core_properties.author
    assert not presentation.core_properties.title


@pytest.mark.leak
def test_verification_passes(source: Path):
    result = l1_scrubber().scrub_file(source)
    assert result.verify.ok, result.verify.summary()


@pytest.mark.leak
def test_names_inside_the_embedded_workbook_do_not_survive(source: Path):
    """The leak this milestone exists for: a zip inside a zip.

    L1 cannot see a name, so this needs the full stack. It regressed once — the
    edits reached the nested elements and were then discarded, because the part
    name handed to the nested package to mark dirty had been prefixed for
    readability and no longer matched anything there.
    """
    pytest.importorskip("gliner2", reason="the 'ner' extra is required")
    output = Scrubber(registry=build_registry("all")).scrub_file(source).output
    for value in ("Mario Rossi", "Anna Bianchi"):
        assert survives_anywhere(output, value) is None, value


def test_the_embedded_workbook_scrub_is_reported(source: Path):
    """Its metadata is gone, which the user should not have to discover."""
    warnings = l1_scrubber().scrub_file(source).warnings
    assert any("embedded at" in warning for warning in warnings)


# -- hidden content -----------------------------------------------------------


def test_hidden_content_is_inventoried(source: Path):
    kinds = {item.kind for item in l1_scrubber().scrub_file(source).hidden_removed}
    assert {"page-thumbnail", "metadata", "offscreen-shape"} <= kinds


def test_a_hidden_slide_is_reported(tmp_path: Path):
    path = build_deck(tmp_path / "nascosta.pptx", chart=False, hidden_slide=True)
    items = [
        item
        for item in l1_scrubber().scrub_file(path).hidden_removed
        if item.kind == "hidden-slide"
    ]
    assert items and "skipped when presenting" in items[0].detail


def test_a_hidden_slide_is_kept(tmp_path: Path):
    """Deleting a slide renumbers the deck and breaks what points at it."""
    path = build_deck(tmp_path / "nascosta.pptx", chart=False, hidden_slide=True)
    output = l1_scrubber().scrub_file(path).output
    assert len(pptx_module.Presentation(io.BytesIO(output)).slides) == 2


def test_an_on_canvas_deck_reports_no_offscreen_shapes(tmp_path: Path):
    path = build_deck(tmp_path / "pulito.pptx", chart=False, off_canvas=False)
    kinds = {item.kind for item in l1_scrubber().scrub_file(path).hidden_removed}
    assert "offscreen-shape" not in kinds
